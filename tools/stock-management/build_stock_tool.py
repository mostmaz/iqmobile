#!/usr/bin/env python3
"""Build the Stock Management workbook from a warehouse export.

Input workbook (one file, three sheets):
  * Inventory      - stock per product (rows) per warehouse tag (columns),
                     tags in row 2, products from row 5 (system export layout)
  * Warehouse tags - tag -> city mapping (any sheet whose name contains "Tag")
  * Sales          - product -> last 30 days sales (any sheet whose name
                     contains "Sales")

Output workbook: every number is an Excel formula over the input sheets, so
quantities, sales, the tag->city mapping and the thresholds on the
Assumptions sheet can be edited in Excel and everything recalculates.
Re-run this script when new products or new warehouse columns are added.

Usage:
    python3 build_stock_tool.py input.xlsx output.xlsx
"""
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

NON_SELLING_CITIES = {"china", "dubai"}   # transit / overseas locations

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_RED = PatternFill("solid", fgColor="F8CBAD")
FILL_DRED = PatternFill("solid", fgColor="FF9999")
FILL_ORANGE = PatternFill("solid", fgColor="FFE699")
FILL_GRAY = PatternFill("solid", fgColor="D9D9D9")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ----------------------------------------------------------------- parsing
def norm(s):
    return " ".join(str(s).split())


def product_key(name):
    """Same rule as the Excel key formula: text inside [..], else whole name.
    Prefixed with '#' so numeric-looking ids stay text in Excel lookups."""
    n = norm(name)
    m = re.match(r"^\[([^\]]*)\]", n)
    return "#" + (m.group(1) if m else n)


def parse_product(name):
    n = norm(name)
    m = re.match(r"^\[([^\]]*)\]\s*(.*)$", n)
    rest = m.group(2) if m else n
    base, variant = rest, ""
    if "(" in rest:
        base = rest.split("(", 1)[0].strip()
        variant = rest.split("(", 1)[1].rstrip(")").strip()
    parts = base.split()
    brand = parts[0].upper() if parts else "OTHER"
    if re.fullmatch(r"[\d.,]+", brand):
        brand = "OTHER"            # junk rows such as "75000 per point on ..."
    model = " ".join(parts[1:]) or "(no model)"
    return brand, model, variant


def read_inputs(path):
    wb = load_workbook(path, data_only=True)
    inv = wb["Inventory"]
    tags_ws = next(ws for ws in wb.worksheets if "tag" in ws.title.lower())
    sales_ws = next(ws for ws in wb.worksheets if "sale" in ws.title.lower())

    inv_rows = [list(r) for r in inv.iter_rows(values_only=True)]
    hdr = inv_rows[1]
    wh_cols = [i for i in range(1, len(hdr)) if hdr[i] not in (None, "")]
    first_wc, last_wc = wh_cols[0] + 1, wh_cols[-1] + 1          # 1-based
    inv_last = max(i + 1 for i, r in enumerate(inv_rows) if r and r[0] not in (None, ""))
    products = OrderedDict()
    for i in range(4, inv_last):
        name = inv_rows[i][0]
        if name in (None, ""):
            continue
        k = product_key(name)
        if k.upper() in {x.upper() for x in products}:
            print(f"WARNING: duplicate product key in Inventory: {name!r}", file=sys.stderr)
            continue
        products[k] = norm(name)

    tags = [(norm(r[0]), norm(r[1])) for r in tags_ws.iter_rows(min_row=2, values_only=True)
            if r[0] not in (None, "")]

    sales_rows = [list(r[:2]) for r in sales_ws.iter_rows(values_only=True)]
    sales_last = max(i + 1 for i, r in enumerate(sales_rows) if r and r[0] not in (None, ""))
    sales_only = OrderedDict()
    for r in sales_rows[1:sales_last]:
        name = r[0]
        if name in (None, "") or norm(name).lower() == "grand total":
            continue
        k = product_key(name)
        if k.upper() not in {x.upper() for x in products} and k not in sales_only:
            sales_only[k] = norm(name)
    for k in list(products) + list(sales_only):
        if any(c in k for c in "*?~"):
            print(f"WARNING: key contains a wildcard character, lookups may misbehave: {k}", file=sys.stderr)

    return dict(inv_rows=inv_rows, inv_last=inv_last, first_wc=first_wc, last_wc=last_wc,
                tags=tags, sales_rows=sales_rows[:sales_last], sales_last=sales_last,
                products=products, sales_only=sales_only)


# ----------------------------------------------------------------- helpers
def style_header(ws, row, c1, c2, fill=FILL_HDR, font=F_HDR):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = font, fill, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, mapping):
    for letter, w in mapping.items():
        ws.column_dimensions[letter].width = w


def apply_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name != FONT:
                cell.font = Font(name=FONT, size=cell.font.size or 10, bold=cell.font.bold,
                                 italic=cell.font.italic, color=cell.font.color)


def status_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OUT OF STOCK"'], fill=FILL_DRED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SHORTAGE"'], fill=FILL_RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SLOW MOVING"'], fill=FILL_ORANGE))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NOT MOVING"'], fill=FILL_GRAY))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_GREEN))


# ------------------------------------------------------------------- build
def build(inp, out_path):
    d = read_inputs(inp)
    wb = Workbook()
    wb.remove(wb.active)

    inv_last, wc0, wc1 = d["inv_last"], col(d["first_wc"]), col(d["last_wc"])
    tags_n = len(d["tags"]) + 1
    sal_n = d["sales_last"]

    cities = sorted({c for _, c in d["tags"]}, key=str.lower)
    n_city = len(cities)
    A_C0, A_C1 = 11, 10 + n_city                       # Assumptions city table rows
    city_rng = lambda c: f"Assumptions!${c}${A_C0}:${c}${A_C1}"

    all_products = list(d["products"].items()) + list(d["sales_only"].items())
    master = []                                        # key, name, brand, model, variant
    for k, name in all_products:
        b, m, v = parse_product(name)
        master.append((k, name, b, m, v))
    master.sort(key=lambda t: (t[2], t[3].lower(), t[1].lower()))
    pm_n = len(master) + 1
    models = sorted({(b, m) for _, _, b, m, _ in master}, key=lambda t: (t[0], t[1].lower()))
    brands = sorted({b for b, _ in models})

    # ---------------- README (filled at the end, placed first)
    ws_readme = wb.create_sheet("README")

    # ---------------- Shortage Summary / WOS by Brand / Model Summary placeholders (order)
    ws_short = wb.create_sheet("Shortage Summary")
    ws_wos = wb.create_sheet("WOS by Brand")
    ws_msum = wb.create_sheet("Model Summary")
    ws_cxm = wb.create_sheet("City x Model")
    ws_as = wb.create_sheet("Assumptions")
    ws_inv = wb.create_sheet("Inventory")
    ws_tags = wb.create_sheet("Warehouse Tags")
    ws_sales = wb.create_sheet("Sales 30 Days")
    ws_pm = wb.create_sheet("Product Master")
    ws_map = wb.create_sheet("Calc Map")
    ws_keys = wb.create_sheet("Calc Keys")
    for ws in (ws_as, ws_inv, ws_tags, ws_sales):
        ws.sheet_properties.tabColor = "FFC000"
    for ws in (ws_short, ws_wos, ws_msum, ws_cxm):
        ws.sheet_properties.tabColor = "70AD47"
    for ws in (ws_pm, ws_map, ws_keys):
        ws.sheet_properties.tabColor = "A6A6A6"

    # ---------------- Inventory (verbatim copy of the export)
    for r, row in enumerate(d["inv_rows"], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws_inv.cell(row=r, column=c, value=v)
    style_header(ws_inv, 2, 1, d["last_wc"])
    ws_inv["A1"] = "Paste the stock export here. Keep the layout: warehouse tags in row 2, product names in column A from row 5, quantities under each tag."
    ws_inv["A1"].font = F_NOTE
    ws_inv.freeze_panes = "B5"
    widths(ws_inv, {"A": 60})

    # ---------------- Warehouse Tags (input)
    ws_tags.append(["Warehouse Tag", "City", "Selling location? (Yes/No)"])
    for tag, city in d["tags"]:
        selling = "No" if city.lower() in NON_SELLING_CITIES else "Yes"
        ws_tags.append([tag, city, selling])
    style_header(ws_tags, 1, 1, 3)
    for r in range(2, tags_n + 1):
        for c in (1, 2, 3):
            ws_tags.cell(row=r, column=c).font = F_INPUT
    ws_tags.cell(row=1, column=3).comment = Comment(
        "Yes = stock here is available to sell in that city. No = in transit / overseas: "
        "excluded from city shortage checks, shown separately as 'in transit'. "
        "Default: China and Dubai = No (assumption, edit as needed).", "Stock tool")
    ws_tags.freeze_panes = "A2"
    widths(ws_tags, {"A": 32, "B": 18, "C": 26})

    # ---------------- Sales 30 Days (input, verbatim copy)
    for row in d["sales_rows"]:
        ws_sales.append(row)
    style_header(ws_sales, 1, 1, 2)
    ws_sales.freeze_panes = "A2"
    widths(ws_sales, {"A": 60, "B": 20})

    # ---------------- Calc Keys (row aligned helper keys)
    ws_keys["A1"] = "Inventory key (same row as Inventory sheet)"
    ws_keys["C1"] = "Sales key (same row as Sales sheet)"
    ws_keys["D1"] = "Sales qty"
    for r in range(5, inv_last + 1):
        t = f"TRIM(Inventory!A{r})"
        ws_keys[f"A{r}"] = f'="#"&IFERROR(IF(LEFT({t},1)="[",MID({t},2,FIND("]",{t})-2),{t}),{t})'
    for r in range(2, sal_n + 1):
        t = f"TRIM('Sales 30 Days'!A{r})"
        ws_keys[f"C{r}"] = f'="#"&IFERROR(IF(LEFT({t},1)="[",MID({t},2,FIND("]",{t})-2),{t}),{t})'
        ws_keys[f"D{r}"] = f"=N('Sales 30 Days'!B{r})"
    style_header(ws_keys, 1, 1, 4)
    widths(ws_keys, {"A": 40, "C": 40, "D": 12})

    # ---------------- Calc Map (warehouse column -> city / selling flag)
    ws_map["A1"], ws_map["A2"], ws_map["A3"] = "Warehouse tag", "City", "Selling?"
    for c in range(d["first_wc"], d["last_wc"] + 1):
        L = col(c)
        ws_map[f"{L}1"] = f"=Inventory!{L}2"
        ws_map[f"{L}2"] = (f"=IFERROR(INDEX('Warehouse Tags'!$B$2:$B${tags_n},"
                           f"MATCH({L}1,'Warehouse Tags'!$A$2:$A${tags_n},0)),\"UNMAPPED\")")
        ws_map[f"{L}3"] = (f"=IFERROR(INDEX('Warehouse Tags'!$C$2:$C${tags_n},"
                           f"MATCH({L}1,'Warehouse Tags'!$A$2:$A${tags_n},0)),\"No\")")
    ws_map["A5"] = "Unmapped tags:"
    ws_map["B5"] = f'=COUNTIF({wc0}2:{wc1}2,"UNMAPPED")'
    style_header(ws_map, 1, 1, d["last_wc"])
    widths(ws_map, {"A": 16})

    # ---------------- Product Master
    pm_hdr = ["Key", "Product", "Brand", "Model", "Brand|Model", "Variant", "Inventory row",
              "Total stock (all warehouses)", "Stock in transit / non-selling", "Stock on hand (selling)",
              "Sales 30 days", "Weekly sales", "WOS (on hand)"] + cities
    ws_pm.append(pm_hdr)
    PC0 = 14                                   # first city column index
    PC1 = PC0 + n_city - 1
    pc0, pc1 = col(PC0), col(PC1)
    inv_rng = f"Inventory!${wc0}$5:${wc1}${inv_last}"
    for i, (k, name, b, m, v) in enumerate(master, start=2):
        ws_pm.cell(row=i, column=1, value=k)
        ws_pm.cell(row=i, column=2, value=name)
        ws_pm.cell(row=i, column=3, value=b)
        ws_pm.cell(row=i, column=4, value=m)
        ws_pm.cell(row=i, column=5, value=f'=C{i}&"|"&D{i}')
        ws_pm.cell(row=i, column=6, value=v)
        ws_pm.cell(row=i, column=7, value=f"=IFERROR(MATCH($A{i},'Calc Keys'!$A$5:$A${inv_last},0),0)")
        ws_pm.cell(row=i, column=8, value=f"=IF($G{i}=0,0,SUM(INDEX({inv_rng},$G{i},0)))")
        ws_pm.cell(row=i, column=9, value=(f"=IF($G{i}=0,0,SUMPRODUCT(('Calc Map'!${wc0}$3:${wc1}$3<>\"Yes\")"
                                           f"*INDEX({inv_rng},$G{i},0)))"))
        ws_pm.cell(row=i, column=10, value=f"=H{i}-I{i}")
        ws_pm.cell(row=i, column=11, value=f"=SUMIF('Calc Keys'!$C$2:$C${sal_n},$A{i},'Calc Keys'!$D$2:$D${sal_n})")
        ws_pm.cell(row=i, column=12, value=f"=MAX(0,K{i})/Assumptions!$B$3*7")
        ws_pm.cell(row=i, column=13, value=f'=IF(L{i}>0,J{i}/L{i},"n/a")')
        for j in range(n_city):
            L = col(PC0 + j)
            ws_pm.cell(row=i, column=PC0 + j,
                       value=f"=IF($G{i}=0,0,SUMPRODUCT(('Calc Map'!${wc0}$2:${wc1}$2={L}$1)*INDEX({inv_rng},$G{i},0)))")
    for r in range(2, pm_n + 1):
        for c in list(range(8, 12)) + list(range(PC0, PC1 + 1)):
            ws_pm.cell(row=r, column=c).number_format = "#,##0"
        ws_pm.cell(row=r, column=12).number_format = "0.0"
        ws_pm.cell(row=r, column=13).number_format = "0.0"
    style_header(ws_pm, 1, 1, PC1)
    ws_pm.freeze_panes = "C2"
    ws_pm.auto_filter.ref = f"A1:{pc1}{pm_n}"
    widths(ws_pm, {"A": 22, "B": 55, "C": 12, "D": 28, "E": 34, "F": 28, "G": 10, "H": 12, "I": 12,
                   "J": 12, "K": 11, "L": 10, "M": 10})
    ws_pm.row_dimensions[1].height = 42

    # ---------------- Assumptions
    ws_as["A1"] = "Assumptions & thresholds"
    ws_as["A1"].font = F_TITLE
    ws_as["A2"] = "Blue text on yellow = input you can change. Everything else is a formula."
    ws_as["A2"].font = F_NOTE
    inputs = [
        (3, "Sales period covered by the Sales sheet (days)", 30,
         "The Sales sheet holds the last 30 days of sales (from the uploaded file). Weekly sales = sales / days * 7."),
        (4, "Shortage threshold: flag a city when weeks of stock (WOS) is below", 4,
         "Assumption: 4 weeks of cover is the minimum a city should hold. Change to your replenishment lead time."),
        (5, "Target cover (weeks) used to size the suggested transfer quantity", 6,
         "Suggested transfer = target cover x city weekly demand - city stock."),
        (6, "Slow-moving threshold: flag when weeks of stock (WOS) is above", 12,
         "Assumption: more than 12 weeks of stock in a city is slow moving."),
        (7, "Ignore a city shortage when expected demand over the target cover is below (units)", 1,
         "Assumption: if a city is expected to sell less than 1 unit of a model within the target cover period, it is 'LOW DEMAND', not a shortage."),
    ]
    for r, label, val, note in inputs:
        ws_as[f"A{r}"] = label
        ws_as[f"B{r}"] = val
        ws_as[f"B{r}"].font, ws_as[f"B{r}"].fill, ws_as[f"B{r}"].border = F_INPUT, FILL_INPUT, BORDER
        ws_as[f"C{r}"] = note
        ws_as[f"C{r}"].font = F_NOTE
    ws_as["A8"] = "Data checks"
    ws_as["A8"].font = F_BOLD
    ws_as["C8"] = "=\"Warehouse tags with no city: \"&'Calc Map'!B5&\"   |   Sold products with no inventory row: \"&COUNTIF('Product Master'!$G$2:$G$" + str(pm_n) + ",0)"
    ws_as["C8"].font = F_NOTE

    hdr = ["City", "Selling location?", "Stock on hand", "Default demand share",
           "Demand share override (optional)", "Demand share used"]
    for c, h in enumerate(hdr, start=1):
        ws_as.cell(row=10, column=c, value=h)
    style_header(ws_as, 10, 1, 6)
    ws_as.cell(row=10, column=4).comment = Comment(
        "ASSUMPTION: sales are only known nationally, not per city. Each city's demand is estimated as "
        "national sales x this share. Default share = the city's share of stock on hand across selling "
        "cities. Type a percentage in the override column to use your own split (e.g. from POS data).",
        "Stock tool")
    for i, city in enumerate(cities):
        r = A_C0 + i
        ws_as.cell(row=r, column=1, value=city)
        ws_as.cell(row=r, column=2, value=(f"=IF(COUNTIFS('Warehouse Tags'!$B$2:$B${tags_n},A{r},"
                                            f"'Warehouse Tags'!$C$2:$C${tags_n},\"Yes\")>0,\"Yes\",\"No\")"))
        L = col(PC0 + i)
        ws_as.cell(row=r, column=3, value=f"=IF(B{r}=\"Yes\",SUM('Product Master'!${L}$2:${L}${pm_n}),0)")
        ws_as.cell(row=r, column=4, value=f"=IF(B{r}=\"Yes\",IF($C${A_C1+1}>0,C{r}/$C${A_C1+1},0),0)")
        ov = ws_as.cell(row=r, column=5)
        ov.font, ov.fill, ov.border = F_INPUT, FILL_INPUT, BORDER
        ws_as.cell(row=r, column=6, value=f"=IF(B{r}=\"Yes\",IF(E{r}<>\"\",E{r},D{r}),0)")
        ws_as.cell(row=r, column=3).number_format = "#,##0"
        for c in (4, 5, 6):
            ws_as.cell(row=r, column=c).number_format = "0.0%"
    tr = A_C1 + 1
    ws_as.cell(row=tr, column=1, value="Total (selling cities)").font = F_BOLD
    ws_as.cell(row=tr, column=3, value=f"=SUM(C{A_C0}:C{A_C1})").number_format = "#,##0"
    ws_as.cell(row=tr, column=4, value=f"=SUM(D{A_C0}:D{A_C1})").number_format = "0.0%"
    ws_as.cell(row=tr, column=6, value=f"=SUM(F{A_C0}:F{A_C1})").number_format = "0.0%"
    ws_as.cell(row=tr + 1, column=1, value="Demand shares used should total 100%. If you override, make sure the 'used' column still sums to 100%.").font = F_NOTE
    widths(ws_as, {"A": 62, "B": 18, "C": 16, "D": 20, "E": 26, "F": 18})

    # ---------------- City x Model
    cx_hdr = ["Brand", "Model", "City", "City stock", "Model sales 30 days (all cities)",
              "City demand share", "City weekly demand (est.)", "City WOS", "Status",
              "Suggested transfer qty", "Model stock on hand (all selling cities)", "Key", "Selling city?"]
    ws_cxm.append(cx_hdr)
    pm_key = f"'Product Master'!$E$2:$E${pm_n}"
    r = 2
    for b, m in models:
        for city in cities:
            ws_cxm.cell(row=r, column=1, value=b)
            ws_cxm.cell(row=r, column=2, value=m)
            ws_cxm.cell(row=r, column=3, value=city)
            ws_cxm.cell(row=r, column=4, value=(
                f"=SUMPRODUCT(({pm_key}=$L{r})*INDEX('Product Master'!${pc0}$2:${pc1}${pm_n},0,"
                f"MATCH($C{r},'Product Master'!${pc0}$1:${pc1}$1,0)))"))
            ws_cxm.cell(row=r, column=5, value=f"=SUMIF({pm_key},$L{r},'Product Master'!$K$2:$K${pm_n})")
            ws_cxm.cell(row=r, column=6, value=f"=IFERROR(INDEX({city_rng('F')},MATCH($C{r},{city_rng('A')},0)),0)")
            ws_cxm.cell(row=r, column=7, value=f"=MAX(0,E{r})/Assumptions!$B$3*7*F{r}")
            ws_cxm.cell(row=r, column=8, value=f'=IF(G{r}>0,D{r}/G{r},"")')
            ws_cxm.cell(row=r, column=9, value=(
                f'=IF(M{r}<>"Yes","NON-SELLING LOCATION",IF(AND(D{r}<=0,E{r}<=0),"NO STOCK / NO SALES",'
                f'IF(E{r}<=0,"NOT MOVING",IF(G{r}<=0,"NO DEMAND SHARE",'
                f'IF(OR(D{r}<=0,H{r}<Assumptions!$B$4),IF(G{r}*Assumptions!$B$5<Assumptions!$B$7,"LOW DEMAND",'
                f'IF(D{r}<=0,"OUT OF STOCK","SHORTAGE")),IF(H{r}>Assumptions!$B$6,"SLOW MOVING","OK"))))))'))
            ws_cxm.cell(row=r, column=10, value=(
                f'=IF(OR(I{r}="SHORTAGE",I{r}="OUT OF STOCK"),MAX(0,ROUNDUP(Assumptions!$B$5*G{r}-D{r},0)),0)'))
            ws_cxm.cell(row=r, column=11, value=f"=SUMIF({pm_key},$L{r},'Product Master'!$J$2:$J${pm_n})")
            ws_cxm.cell(row=r, column=12, value=f'=A{r}&"|"&B{r}')
            ws_cxm.cell(row=r, column=13, value=f"=IFERROR(INDEX({city_rng('B')},MATCH($C{r},{city_rng('A')},0)),\"No\")")
            for c in (4, 5, 10, 11):
                ws_cxm.cell(row=r, column=c).number_format = "#,##0"
            ws_cxm.cell(row=r, column=6).number_format = "0.0%"
            ws_cxm.cell(row=r, column=7).number_format = "0.0"
            ws_cxm.cell(row=r, column=8).number_format = "0.0"
            r += 1
    cx_n = r - 1
    style_header(ws_cxm, 1, 1, 13)
    ws_cxm.freeze_panes = "D2"
    ws_cxm.auto_filter.ref = f"A1:M{cx_n}"
    status_cf(ws_cxm, f"I2:I{cx_n}")
    widths(ws_cxm, {"A": 12, "B": 30, "C": 15, "D": 10, "E": 14, "F": 11, "G": 12, "H": 9, "I": 22,
                    "J": 12, "K": 14, "L": 34, "M": 10})
    ws_cxm.row_dimensions[1].height = 45
    ws_cxm.cell(row=1, column=9).comment = Comment(
        "OUT OF STOCK: model sells but this city has 0 stock. SHORTAGE: city WOS below the shortage threshold. "
        "OK: between thresholds. SLOW MOVING: city WOS above the slow-moving threshold. "
        "NOT MOVING: stock in this city but zero sales in the whole sales period. "
        "LOW DEMAND: city would be short but is expected to sell less than the minimum units in the target cover period. "
        "NO STOCK / NO SALES: nothing to act on. NON-SELLING LOCATION: transit/overseas city.", "Stock tool")

    cxA = f"'City x Model'!$A$2:$A${cx_n}"
    cxC = f"'City x Model'!$C$2:$C${cx_n}"
    cxD = f"'City x Model'!$D$2:$D${cx_n}"
    cxI = f"'City x Model'!$I$2:$I${cx_n}"
    cxJ = f"'City x Model'!$J$2:$J${cx_n}"
    cxL = f"'City x Model'!$L$2:$L${cx_n}"

    # ---------------- WOS by Brand
    ws_wos["A1"] = "Weeks of Sales (WOS) per brand"
    ws_wos["A1"].font = F_TITLE
    ws_wos["A2"] = "WOS = stock / weekly sales. Weekly sales = sales in period / period days x 7 (see Assumptions)."
    ws_wos["A2"].font = F_NOTE
    hdr = ["Brand", "Stock on hand (selling cities)", "Stock in transit / non-selling", "Total stock",
           "Sales 30 days", "Weekly sales", "WOS (on hand)", "WOS (incl. transit)", "# SKUs",
           "# SKUs sold out (sales > 0, stock 0)", "# SKUs not moving (stock > 0, sales 0)", "Brand status"]
    for c, h in enumerate(hdr, start=1):
        ws_wos.cell(row=4, column=c, value=h)
    style_header(ws_wos, 4, 1, 12)
    pmC, pmJ, pmI, pmK = (f"'Product Master'!${x}$2:${x}${pm_n}" for x in "CJIK")
    for i, b in enumerate(brands, start=5):
        ws_wos.cell(row=i, column=1, value=b)
        ws_wos.cell(row=i, column=2, value=f"=SUMIF({pmC},$A{i},{pmJ})")
        ws_wos.cell(row=i, column=3, value=f"=SUMIF({pmC},$A{i},{pmI})")
        ws_wos.cell(row=i, column=4, value=f"=B{i}+C{i}")
        ws_wos.cell(row=i, column=5, value=f"=SUMIF({pmC},$A{i},{pmK})")
        ws_wos.cell(row=i, column=6, value=f"=MAX(0,E{i})/Assumptions!$B$3*7")
        ws_wos.cell(row=i, column=7, value=f'=IF(F{i}>0,B{i}/F{i},"n/a")')
        ws_wos.cell(row=i, column=8, value=f'=IF(F{i}>0,D{i}/F{i},"n/a")')
        ws_wos.cell(row=i, column=9, value=f"=COUNTIF({pmC},$A{i})")
        ws_wos.cell(row=i, column=10, value=f'=COUNTIFS({pmC},$A{i},{pmJ},"<=0",{pmK},">0")')
        ws_wos.cell(row=i, column=11, value=f'=COUNTIFS({pmC},$A{i},{pmJ},">0",{pmK},"<=0")')
        ws_wos.cell(row=i, column=12, value=(
            f'=IF(F{i}=0,"NO SALES",IF(G{i}<Assumptions!$B$4,"SHORTAGE",IF(G{i}>Assumptions!$B$6,"SLOW MOVING","OK")))'))
        for c in (2, 3, 4, 5, 9, 10, 11):
            ws_wos.cell(row=i, column=c).number_format = "#,##0"
        for c in (6, 7, 8):
            ws_wos.cell(row=i, column=c).number_format = "0.0"
    tr = 5 + len(brands)
    ws_wos.cell(row=tr, column=1, value="TOTAL").font = F_BOLD
    for c in (2, 3, 4, 5, 6, 9, 10, 11):
        L = col(c)
        cell = ws_wos.cell(row=tr, column=c, value=f"=SUM({L}5:{L}{tr-1})")
        cell.font, cell.number_format = F_BOLD, "#,##0" if c != 6 else "0.0"
    ws_wos.cell(row=tr, column=7, value=f'=IF(F{tr}>0,B{tr}/F{tr},"n/a")').number_format = "0.0"
    ws_wos.cell(row=tr, column=8, value=f'=IF(F{tr}>0,D{tr}/F{tr},"n/a")').number_format = "0.0"
    status_cf(ws_wos, f"L5:L{tr-1}")
    ws_wos.conditional_formatting.add(f"G5:G{tr}", CellIsRule(operator="lessThan", formula=["Assumptions!$B$4"], fill=FILL_RED))
    ws_wos.conditional_formatting.add(f"G5:G{tr}", CellIsRule(operator="greaterThan", formula=["Assumptions!$B$6"], fill=FILL_ORANGE))
    ws_wos.freeze_panes = "B5"
    widths(ws_wos, {"A": 14, "B": 16, "C": 16, "D": 12, "E": 12, "F": 11, "G": 11, "H": 12, "I": 9,
                    "J": 16, "K": 16, "L": 14})
    ws_wos.row_dimensions[4].height = 45

    # ---------------- Model Summary (per brand per model, national)
    ws_msum["A1"] = "Model summary (per brand, per model, all cities)"
    ws_msum["A1"].font = F_TITLE
    ws_msum["A2"] = "Use the filters: sort by 'Cities short' to see which models need transfers, by 'Cities not moving' for dead stock."
    ws_msum["A2"].font = F_NOTE
    hdr = ["Brand", "Model", "Key", "Stock on hand (selling cities)", "Stock in transit", "Sales 30 days",
           "Weekly sales", "WOS (on hand)", "# SKUs", "Cities short / out of stock", "Cities not moving",
           "Units to transfer (all cities)", "Model status"]
    for c, h in enumerate(hdr, start=1):
        ws_msum.cell(row=4, column=c, value=h)
    style_header(ws_msum, 4, 1, 13)
    pmE = f"'Product Master'!$E$2:$E${pm_n}"
    for i, (b, m) in enumerate(models, start=5):
        ws_msum.cell(row=i, column=1, value=b)
        ws_msum.cell(row=i, column=2, value=m)
        ws_msum.cell(row=i, column=3, value=f'=A{i}&"|"&B{i}')
        ws_msum.cell(row=i, column=4, value=f"=SUMIF({pmE},$C{i},{pmJ})")
        ws_msum.cell(row=i, column=5, value=f"=SUMIF({pmE},$C{i},{pmI})")
        ws_msum.cell(row=i, column=6, value=f"=SUMIF({pmE},$C{i},{pmK})")
        ws_msum.cell(row=i, column=7, value=f"=MAX(0,F{i})/Assumptions!$B$3*7")
        ws_msum.cell(row=i, column=8, value=f'=IF(G{i}>0,D{i}/G{i},"n/a")')
        ws_msum.cell(row=i, column=9, value=f"=COUNTIF({pmE},$C{i})")
        ws_msum.cell(row=i, column=10, value=f'=COUNTIFS({cxL},$C{i},{cxI},"SHORTAGE")+COUNTIFS({cxL},$C{i},{cxI},"OUT OF STOCK")')
        ws_msum.cell(row=i, column=11, value=f'=COUNTIFS({cxL},$C{i},{cxI},"NOT MOVING")')
        ws_msum.cell(row=i, column=12, value=f"=SUMIFS({cxJ},{cxL},$C{i})")
        ws_msum.cell(row=i, column=13, value=(
            f'=IF(AND(D{i}<=0,F{i}<=0),"NO STOCK / NO SALES",IF(F{i}<=0,"NOT MOVING",IF(D{i}<=0,"OUT OF STOCK",'
            f'IF(H{i}<Assumptions!$B$4,"SHORTAGE",IF(H{i}>Assumptions!$B$6,"SLOW MOVING","OK")))))'))
        for c in (4, 5, 6, 9, 10, 11, 12):
            ws_msum.cell(row=i, column=c).number_format = "#,##0"
        for c in (7, 8):
            ws_msum.cell(row=i, column=c).number_format = "0.0"
    ms_n = 4 + len(models)
    status_cf(ws_msum, f"M5:M{ms_n}")
    ws_msum.freeze_panes = "C5"
    ws_msum.auto_filter.ref = f"A4:M{ms_n}"
    widths(ws_msum, {"A": 12, "B": 32, "C": 34, "D": 14, "E": 12, "F": 11, "G": 10, "H": 10, "I": 8,
                     "J": 14, "K": 12, "L": 14, "M": 20})
    ws_msum.row_dimensions[4].height = 45

    # ---------------- Shortage Summary (city x brand matrices)
    ws = ws_short
    ws["A1"] = "Shortage & slow-stock summary by city and brand"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Counts are models (brand + model, all colours/memory variants combined). "
                "Open 'City x Model' and filter by Status to see the exact models behind any number.")
    ws["A2"].font = F_NOTE
    nb = len(brands)
    wos_brand_row = {b: 5 + i for i, b in enumerate(brands)}

    def matrix(top, title, cell_formula, fmt="#,##0", cf=None):
        ws.cell(row=top, column=1, value=title).font = F_BOLD
        ws.cell(row=top + 1, column=1, value="City")
        ws.cell(row=top + 1, column=2, value="Selling?")
        for j, b in enumerate(brands):
            ws.cell(row=top + 1, column=3 + j, value=b)
        ws.cell(row=top + 1, column=3 + nb, value="Total")
        style_header(ws, top + 1, 1, 3 + nb)
        for i, city in enumerate(cities):
            rr = top + 2 + i
            ws.cell(row=rr, column=1, value=city)
            ws.cell(row=rr, column=2, value=f"=IFERROR(INDEX({city_rng('B')},MATCH($A{rr},{city_rng('A')},0)),\"No\")")
            for j, b in enumerate(brands):
                L = col(3 + j)
                cell = ws.cell(row=rr, column=3 + j, value=cell_formula(rr, L, b))
                cell.number_format = fmt
            tot = ws.cell(row=rr, column=3 + nb, value=f"=SUM(C{rr}:{col(2+nb)}{rr})")
            tot.number_format, tot.font = fmt, F_BOLD
        rr = top + 2 + n_city
        ws.cell(row=rr, column=1, value="Total").font = F_BOLD
        for j in range(nb + 1):
            L = col(3 + j)
            cell = ws.cell(row=rr, column=3 + j, value=f"=SUM({L}{top+2}:{L}{rr-1})")
            cell.number_format, cell.font = fmt, F_BOLD
        rng = f"C{top+2}:{col(2+nb)}{rr-1}"
        if cf:
            cf(rng)
        return rr + 2

    def cf_pos(fill):
        return lambda rng: ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=fill))

    top = 4
    top = matrix(top, "1) Models facing SHORTAGE or OUT OF STOCK (count of models) - city x brand",
                 lambda rr, L, b: f'=COUNTIFS({cxA},{L}$5,{cxC},$A{rr},{cxI},"SHORTAGE")+COUNTIFS({cxA},{L}$5,{cxC},$A{rr},{cxI},"OUT OF STOCK")',
                 cf=cf_pos(FILL_RED))
    t2 = top
    top = matrix(top, "2) Suggested units to transfer in (to reach target cover) - city x brand",
                 lambda rr, L, b: f"=SUMIFS({cxJ},{cxA},{L}${t2+1},{cxC},$A{rr})", cf=cf_pos(FILL_ORANGE))
    t3 = top
    top = matrix(top, "3) Models NOT MOVING (stock in city, zero sales in the period) - count of models",
                 lambda rr, L, b: f'=COUNTIFS({cxA},{L}${t3+1},{cxC},$A{rr},{cxI},"NOT MOVING")', cf=cf_pos(FILL_GRAY))
    t4 = top
    top = matrix(top, "4) Units of NOT MOVING stock sitting in the city - city x brand",
                 lambda rr, L, b: f'=SUMIFS({cxD},{cxA},{L}${t4+1},{cxC},$A{rr},{cxI},"NOT MOVING")', cf=cf_pos(FILL_GRAY))
    t5 = top
    top = matrix(top, "5) Models SLOW MOVING (city WOS above threshold) - count of models",
                 lambda rr, L, b: f'=COUNTIFS({cxA},{L}${t5+1},{cxC},$A{rr},{cxI},"SLOW MOVING")', cf=cf_pos(FILL_ORANGE))
    t6 = top

    def wos_cell(rr, L, b):
        wr = wos_brand_row[b]
        stock = f"SUMIFS({cxD},{cxA},{L}${t6+1},{cxC},$A{rr})"
        dem = f"'WOS by Brand'!$F${wr}*IFERROR(INDEX({city_rng('F')},MATCH($A{rr},{city_rng('A')},0)),0)"
        return f'=IF({dem}>0,{stock}/({dem}),"")'
    top = matrix(top, "6) Weeks of stock (WOS) per brand per city = city stock / (brand weekly sales x city demand share)",
                 wos_cell, fmt="0.0",
                 cf=lambda rng: (ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["Assumptions!$B$4"], fill=FILL_RED)),
                                 ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["Assumptions!$B$6"], fill=FILL_ORANGE))))
    # the Total column and total row of the WOS matrix are meaningless; blank them
    for i in range(n_city + 1):
        ws.cell(row=t6 + 2 + i, column=3 + nb).value = None
    ws.cell(row=t6 + 1, column=3 + nb).value = None
    for c in range(2, 4 + nb):
        ws.cell(row=t6 + 2 + n_city, column=c).value = None
    widths(ws, {"A": 20, "B": 9})
    for j in range(nb + 1):
        ws.column_dimensions[col(3 + j)].width = 11
    ws.freeze_panes = "C4"

    # ---------------- README
    ws = ws_readme
    lines = [
        ("Stock Management Tool - shortage, non-moving stock and weeks of sales (WOS)", F_TITLE),
        ("", None),
        ("WHAT IT ANSWERS", F_BOLD),
        ("1. Which city will face a shortage, and for which model (per brand): sheet 'Shortage Summary' (tables 1 & 2) and sheet 'City x Model' filtered on Status = SHORTAGE or OUT OF STOCK.", None),
        ("2. Which city holds products that are not moving (per brand per model): 'Shortage Summary' (tables 3 & 4) and 'City x Model' filtered on Status = NOT MOVING (or SLOW MOVING).", None),
        ("3. Weeks of sales (WOS) per brand: sheet 'WOS by Brand' (national) and 'Shortage Summary' table 6 (per brand per city).", None),
        ("'Model Summary' gives the same per brand per model nationally, with a count of cities in shortage / not moving.", None),
        ("", None),
        ("HOW TO REFRESH WITH NEW DATA (all numbers are formulas and recalculate automatically)", F_BOLD),
        ("- Inventory: paste the new stock export over the 'Inventory' sheet. Keep the layout: warehouse tags in row 2, product names in column A from row 5, quantities under each tag.", None),
        ("- Sales 30 Days: paste the new product / quantity list (column A product, column B quantity).", None),
        ("- Warehouse Tags: add any new tag with its city, and say whether it is a selling location (Yes) or transit/overseas (No).", None),
        ("- Assumptions: set the sales period in days, the shortage threshold, the target cover and the slow-moving threshold. Optionally override each city's demand share.", None),
        ("- If NEW products or NEW warehouse columns were added, re-run tools/stock-management/build_stock_tool.py so the product and city lists are regenerated (the formulas only cover the rows and columns that existed when the file was built).", None),
        ("", None),
        ("DEFINITIONS", F_BOLD),
        ("Product key: the code inside [ ] in the product name (or the full name when there is no code). It links Inventory rows to Sales rows.", None),
        ("Brand = first word of the product name (upper case). Model = the words after the brand up to the first '('. Variant = the text inside '( )' (colour, memory). Models combine all their variants.", None),
        ("Stock on hand = stock in warehouses whose tag is a selling location (Yes). Stock in transit / non-selling = all other warehouses (China, Dubai, transit by default).", None),
        ("Weekly sales = sales in the period / period days x 7.   WOS = stock / weekly sales.", None),
        ("City weekly demand (estimate) = model weekly sales (all cities) x city demand share. ASSUMPTION: the sales file is not split by city, so demand is allocated by the city's share of stock on hand (editable on Assumptions).", None),
        ("City WOS = city stock / city weekly demand.", None),
        ("Status per city per model: OUT OF STOCK = model sells, city has 0 stock | SHORTAGE = city WOS below the shortage threshold | OK | SLOW MOVING = city WOS above the slow-moving threshold | NOT MOVING = city has stock but the model had zero sales in the whole period | LOW DEMAND = would be short, but the city is expected to sell less than the minimum units within the target cover (no action) | NO STOCK / NO SALES = nothing to act on | NON-SELLING LOCATION.", None),
        ("Suggested transfer qty = target cover (weeks) x city weekly demand - city stock, for SHORTAGE / OUT OF STOCK rows only.", None),
        ("Negative sales (returns) are treated as zero demand.", None),
        ("", None),
        ("SHEETS", F_BOLD),
        ("Green tabs = results. Orange tabs = inputs you paste or edit. Grey tabs = helper calculations (do not edit).", None),
        ("Colour legend: red = out of stock / shortage, orange = slow moving, grey = not moving, green = OK. Blue text on yellow = input cell.", None),
        ("", None),
        ("Source: Inventory, warehouse tags and 30-day sales come from the uploaded file 'wearhouse_inputs.xlsx'. Thresholds on the Assumptions sheet are assumptions to adjust.", F_NOTE),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font or F_BASE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    for w in wb.worksheets:
        apply_font(w)
    wb.save(out_path)
    print(f"Built {out_path}: {len(master)} SKUs, {len(models)} models, {len(brands)} brands, "
          f"{n_city} cities, {cx_n-1} city x model rows")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2])
